from unittest.mock import MagicMock, patch

from omnidebate.core.logic import (
    bool_search,
    impact_tree,
    parse_verbatim_docx,
    smart_cite,
    summarize_with_ollama,
)


def test_parse_verbatim_docx(tmp_path):
    from docx import Document

    d_path = tmp_path / "test.docx"
    doc = Document()
    # Add a paragraph with a bold run (heuristic for tag)
    p = doc.add_paragraph()
    p.add_run("Test Tag").bold = True
    doc.add_paragraph("Cite 2024")
    doc.add_paragraph("This is the evidence body text.")
    doc.save(str(d_path))

    cards = parse_verbatim_docx(str(d_path))
    assert len(cards) == 1
    assert cards[0]["tag"] == "Test Tag"
    assert "body text" in cards[0]["text"]


@patch("requests.post")
def test_summarize_with_ollama(mock_post):
    mock_resp = MagicMock()
    mock_resp.json.return_value = {"response": "TAG: AI CLAIM\nBRIEF: - Point 1\n- Point 2"}
    mock_post.return_value = mock_resp

    tag, brief = summarize_with_ollama("Some long debate text", "http://localhost:11434", "llama3")
    assert tag == "AI CLAIM"
    assert "Point 1" in brief


def test_smart_cite():
    assert (
        smart_cite("Smith", "2023", "Trade Wars", "Brookings")
        == 'Smith, (2023), "Trade Wars", — Brookings'
    )
    assert smart_cite("Smith", "2023", "", "") == "Smith, (2023)"
    assert smart_cite("", "", "", "") == "Unknown Source"


def test_impact_tree():
    card_econ = {
        "side": "AFF",
        "tag": "Economy Tag",
        "text": "This card mentions GDP and economy.",
    }
    tree_econ = impact_tree(card_econ)
    assert "Economic Resilience" in tree_econ

    card_social = {
        "side": "NEG",
        "tag": "Social Tag",
        "text": "This card mentions health and poverty.",
    }
    tree_social = impact_tree(card_social)
    assert "Human Capital" in tree_social

    card_none = {"side": "AFF", "tag": "None Tag", "text": "Nothing interesting here."}
    tree_none = impact_tree(card_none)
    assert "❓ Unmapped" in tree_none


def test_bool_search():
    cards = [
        {"id": 1, "tag": "apple", "text": "fruit", "citation": "source1"},
        {"id": 2, "tag": "banana", "text": "fruit", "citation": "source2"},
        {"id": 3, "tag": "carrot", "text": "vegetable", "citation": "source3"},
    ]

    # Simple search
    assert len(bool_search(cards, "apple")) == 1
    assert len(bool_search(cards, "fruit")) == 2

    # AND search
    assert len(bool_search(cards, "apple AND fruit")) == 1
    assert len(bool_search(cards, "apple AND vegetable")) == 0

    # OR search
    assert len(bool_search(cards, "apple OR banana")) == 2

    # NOT search
    assert len(bool_search(cards, "fruit AND NOT apple")) == 1
    assert bool_search(cards, "fruit AND NOT apple")[0]["id"] == 2
