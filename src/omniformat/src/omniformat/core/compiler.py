"""OmniFormat Standalone Document Compiler.

Converts, formats, and compiles general Markdown files into:
1. Prettified Markdown (.md)
2. Beautiful Interactive HTML (.html) with responsive visual styling.
3. Professional Word Documents (.docx) using python-docx.
© 2026 OmniSuite. Built for the Digital Grandmaster.
"""

import datetime
import html
import os
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def parse_markdown_to_html(md_text: str) -> str:
    """A clean, regex-based Markdown to HTML converter supporting lists, tables, and blockquotes with layout indentation."""
    lines = md_text.splitlines()
    html_lines = []
    in_list = False
    in_code_block = False
    code_block_content = []
    code_lang = ""
    current_level = 0

    # Table state
    in_table = False
    table_headers = []
    table_rows = []

    for line in lines:
        # 1. Fenced Code Blocks
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                joined_code = "\n".join(code_block_content)
                indent = (current_level - 1) * 20 if current_level > 1 else 0
                style_attr = f' style="margin-left: {indent}px;"' if indent > 0 else ""
                html_lines.append(
                    f'<pre{style_attr}><code class="language-{code_lang}">{html.escape(joined_code)}</code></pre>'
                )
                code_block_content = []
                code_lang = ""
            else:
                in_code_block = True
                code_lang = line[3:].strip()
            continue

        if in_code_block:
            code_block_content.append(line)
            continue

        # Close lists/tables if needed
        if in_list and not line.strip().startswith("-") and not line.strip().startswith("*"):
            html_lines.append("</ul>")
            in_list = False

        if in_table and ("|" not in line or line.strip() == ""):
            html_lines.append(render_html_table(table_headers, table_rows))
            in_table = False
            table_headers = []
            table_rows = []

        # Empty line
        if not line.strip():
            html_lines.append("<br/>")
            continue

        # 2. Headings
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading_match:
            hashes, content = heading_match.groups()
            level = len(hashes)
            current_level = level
            formatted_content = parse_inline_styles(content)
            indent = (level - 1) * 20 if level > 1 else 0
            style_attr = f' style="margin-left: {indent}px;"' if indent > 0 else ""
            html_lines.append(f"<h{level}{style_attr}>{formatted_content}</h{level}>")
            continue

        # 3. Blockquotes
        if line.strip().startswith(">"):
            content = line.strip()[1:].strip()
            formatted_content = parse_inline_styles(content)
            indent = (current_level - 1) * 20 if current_level > 1 else 0
            style_attr = f' style="margin-left: {indent}px;"' if indent > 0 else ""
            html_lines.append(f"<blockquote{style_attr}>{formatted_content}</blockquote>")
            continue

        # 4. Unordered Lists
        list_match = re.match(r"^\s*[\-\*]\s+(.*)$", line)
        if list_match:
            content = list_match.group(1)
            formatted_content = parse_inline_styles(content)
            if not in_list:
                indent = (current_level - 1) * 20 if current_level > 1 else 0
                style_attr = f' style="margin-left: {indent}px;"' if indent > 0 else ""
                html_lines.append(f"<ul{style_attr}>")
                in_list = True
            html_lines.append(f"<li>{formatted_content}</li>")
            continue

        # 5. Tables
        if "|" in line:
            parts = [p.strip() for p in line.split("|")[1:-1]]
            # Check for divider line like |---|---|
            if all(re.match(r"^:?\-+:?$", p) for p in parts) and parts:
                continue
            if not in_table:
                in_table = True
                table_headers = parts
            else:
                table_rows.append(parts)
            continue

        # 6. Standard Paragraph
        formatted_content = parse_inline_styles(line)
        indent = (current_level - 1) * 20 if current_level > 1 else 0
        style_attr = f' style="margin-left: {indent}px;"' if indent > 0 else ""
        html_lines.append(f"<p{style_attr}>{formatted_content}</p>")

    # Clean up trailing open tags
    if in_list:
        html_lines.append("</ul>")
    if in_table:
        html_lines.append(render_html_table(table_headers, table_rows))

    return "\n".join(html_lines)


def render_html_table(headers: list[str], rows: list[list[str]]) -> str:
    """Generate HTML code for table arrays."""
    t_html = ["<table>"]
    if headers:
        t_html.append("<thead><tr>")
        for h in headers:
            t_html.append(f"<th>{parse_inline_styles(h)}</th>")
        t_html.append("</tr></thead>")
    if rows:
        t_html.append("<tbody>")
        for row in rows:
            t_html.append("<tr>")
            for col in row:
                t_html.append(f"<td>{parse_inline_styles(col)}</td>")
            t_html.append("</tr>")
        t_html.append("</tbody>")
    t_html.append("</table>")
    return "\n".join(t_html)


def parse_inline_styles(text: str) -> str:
    """Format bold, italic, code, and links in HTML runs."""
    text = html.escape(text)

    # Markdown links: [label](url)
    text = re.sub(
        r"\[([^\]]+)\]\(([^)]+)\)",
        r'<a href="\2" target="_blank">\1</a>',
        text,
    )

    # Inline Code: `code`
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)

    # Bold: **bold** or __bold__
    text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"__(.*?)__", r"<strong>\1</strong>", text)

    # Italic: *italic* or _italic_
    text = re.sub(r"\*(.*?)\*", r"<em>\1</em>", text)
    text = re.sub(r"_(.*?)_", r"<em>\1</em>", text)

    return text


def compile_to_html(src_path: str, theme: str = "cyberdark") -> str:
    """Compiles a Markdown file into a beautiful single-page HTML file."""
    with open(src_path, encoding="utf-8") as f:
        md_content = f.read()

    filename = os.path.basename(src_path)
    title = os.path.splitext(filename)[0].upper()
    body_html = parse_markdown_to_html(md_content)
    today = datetime.date.today().strftime("%Y-%m-%d")

    # Define CSS theme styling rules
    if theme == "cyberdark":
        css_styles = """
            :root {
                --bg-color: #020617;
                --surface-color: #0f172a;
                --border-color: #334155;
                --text-primary: #f1f5f9;
                --text-secondary: #94a3b8;
                --accent-color: #10b981;
                --code-bg: #1e293b;
            }
        """
    else:  # warmpaper / light
        css_styles = """
            :root {
                --bg-color: #faf8f5;
                --surface-color: #ffffff;
                --border-color: #e2e8f0;
                --text-primary: #1e293b;
                --text-secondary: #64748b;
                --accent-color: #4f46e5;
                --code-bg: #f1f5f9;
            }
        """

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title} - Styled Document</title>
    <style>
        {css_styles}
        body {{
            background-color: var(--bg-color);
            color: var(--text-primary);
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
            line-height: 1.6;
            margin: 0;
            padding: 0;
            display: flex;
            justify-content: center;
        }}
        #container {{
            width: 100%;
            max-width: 800px;
            padding: 3rem 2rem;
            box-sizing: border-box;
            background: var(--surface-color);
            min-height: 100vh;
            border-left: 1px solid var(--border-color);
            border-right: 1px solid var(--border-color);
        }}
        header {{
            border-bottom: 2px solid var(--accent-color);
            padding-bottom: 1rem;
            margin-bottom: 2rem;
        }}
        h1 {{ font-size: 2.2rem; color: var(--text-primary); margin-top: 0; }}
        h2 {{ font-size: 1.6rem; color: var(--text-primary); margin-top: 2rem; border-bottom: 1px solid var(--border-color); padding-bottom: 0.5rem; }}
        h3 {{ font-size: 1.25rem; color: var(--text-primary); }}
        p {{ margin-bottom: 1.25rem; color: var(--text-secondary); }}
        a {{ color: var(--accent-color); text-decoration: none; }}
        a:hover {{ text-decoration: underline; }}
        ul {{ margin-bottom: 1.25rem; padding-left: 1.5rem; color: var(--text-secondary); }}
        li {{ margin-bottom: 0.5rem; }}
        blockquote {{
            border-left: 4px solid var(--accent-color);
            background: rgba(16, 185, 129, 0.05);
            margin: 1.5rem 0;
            padding: 1rem 1.5rem;
            font-style: italic;
            border-radius: 0 8px 8px 0;
        }}
        pre {{
            background: var(--code-bg);
            border: 1px solid var(--border-color);
            padding: 1rem;
            border-radius: 8px;
            overflow-x: auto;
        }}
        code {{
            font-family: SFMono-Regular, Consolas, Monaco, monospace;
            font-size: 0.9em;
            background: rgba(0,0,0,0.1);
            padding: 0.2em 0.4em;
            border-radius: 4px;
        }}
        pre code {{
            background: none;
            padding: 0;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 1.5rem;
            font-size: 0.95rem;
        }}
        th, td {{
            border: 1px solid var(--border-color);
            padding: 0.75rem 1rem;
            text-align: left;
        }}
        th {{
            background: rgba(0,0,0,0.05);
            font-weight: 600;
        }}
        footer {{
            margin-top: 4rem;
            border-top: 1px solid var(--border-color);
            padding-top: 1rem;
            text-align: center;
            font-size: 0.85rem;
            color: var(--text-secondary);
        }}
    </style>
</head>
<body>
    <div id="container">
        <header>
            <h1>{title}</h1>
            <p>Formatted on: {today} | Source: {filename}</p>
        </header>
        <main>
            {body_html}
        </main>
        <footer>
            Formatted by OmniFormat Standalone Markdown Workstation
        </footer>
    </div>
</body>
</html>
"""
    dest_path = os.path.splitext(src_path)[0] + f"_compiled_{theme}.html"
    with open(dest_path, "w", encoding="utf-8") as out:
        out.write(html_content)
    return dest_path


def compile_to_docx(src_path: str) -> str:
    """Compiles a Markdown file into a styled Word Document (.docx) with indentation support."""
    with open(src_path, encoding="utf-8") as f:
        md_content = f.read()

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    filename = os.path.basename(src_path)
    title_text = os.path.splitext(filename)[0].upper()

    # Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title_text)
    title_run.bold = True
    title_run.font.name = "Trebuchet MS"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(79, 70, 229)  # Deep Indigo

    # Subtitle Info
    today = datetime.date.today().strftime("%Y-%m-%d")
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"Formatted: {today} | Source: {filename}")
    sub_run.italic = True
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(120, 120, 120)

    # Divider
    div_p = doc.add_paragraph()
    div_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div_p.add_run("____________________________________________________").font.color.rgb = RGBColor(
        200, 200, 200
    )

    lines = md_content.splitlines()
    in_code_block = False
    code_block_lines = []
    current_level = 0

    for line in lines:
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_p = doc.add_paragraph()
                indent = 0.4 + (current_level - 1) * 0.25 if current_level > 1 else 0.4
                code_p.paragraph_format.left_indent = Inches(indent)
                code_run = code_p.add_run("\n".join(code_block_lines))
                code_run.font.name = "Consolas"
                code_run.font.size = Pt(9.5)
                code_run.font.color.rgb = RGBColor(80, 80, 80)
                code_block_lines = []
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        if not line.strip():
            continue

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading_match:
            hashes, content = heading_match.groups()
            level = len(hashes)
            current_level = level
            h_p = doc.add_paragraph()
            if level > 1:
                h_p.paragraph_format.left_indent = Inches((level - 1) * 0.25)
            h_run = h_p.add_run(content)
            h_run.bold = True
            h_run.font.name = "Trebuchet MS"
            if level == 1:
                h_run.font.size = Pt(16)
                h_run.font.color.rgb = RGBColor(79, 70, 229)
            elif level == 2:
                h_run.font.size = Pt(14)
                h_run.font.color.rgb = RGBColor(30, 41, 59)
            else:
                h_run.font.size = Pt(12)
                h_run.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Blockquotes
        if line.strip().startswith(">"):
            content = line.strip()[1:].strip()
            bq_p = doc.add_paragraph()
            indent = 0.4 + (current_level - 1) * 0.25 if current_level > 1 else 0.4
            bq_p.paragraph_format.left_indent = Inches(indent)
            bq_run = bq_p.add_run(content)
            bq_run.font.name = "Georgia"
            bq_run.font.size = Pt(10.5)
            bq_run.italic = True
            bq_run.font.color.rgb = RGBColor(100, 116, 139)
            continue

        # Unordered list items
        list_match = re.match(r"^\s*[\-\*]\s+(.*)$", line)
        if list_match:
            content = list_match.group(1)
            li_p = doc.add_paragraph(style="List Bullet")
            indent = 0.2 + (current_level - 1) * 0.25 if current_level > 1 else 0.2
            li_p.paragraph_format.left_indent = Inches(indent)
            apply_inline_docx_formatting(li_p, content)
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        if current_level > 1:
            p.paragraph_format.left_indent = Inches((current_level - 1) * 0.25)
        apply_inline_docx_formatting(p, line)

    dest_path = os.path.splitext(src_path)[0] + "_compiled.docx"
    doc.save(dest_path)
    return dest_path


def apply_inline_docx_formatting(paragraph: Any, text: str) -> None:
    """Helper to parse inline markdown tags (bold/italic/code) and add runs to docx paragraph."""
    # Split text into segments of code, bold, italic, or plain text
    # A simple tokenizer:
    tokens = re.split(r"(\*\*.*?\*\*|__.*?__|`.*?`|\*.*?\*|_.*?_)", text)
    for token in tokens:
        if not token:
            continue

        if token.startswith("`") and token.endswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(180, 50, 50)
        elif (token.startswith("**") and token.endswith("**")) or (
            token.startswith("__") and token.endswith("__")
        ):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(11)
        elif (token.startswith("*") and token.endswith("*")) or (
            token.startswith("_") and token.endswith("_")
        ):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            r.font.name = "Calibri"
            r.font.size = Pt(11)
        else:
            r = paragraph.add_run(token)
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(50, 50, 50)


def prettify_markdown_file(src_path: str) -> str:
    """Rewrites a Markdown file with standardized markdown formatting."""
    with open(src_path, encoding="utf-8") as f:
        md_content = f.read()

    from omniformat.core.compliance import auto_correct_markdown

    corrected = auto_correct_markdown(md_content)
    with open(src_path, "w", encoding="utf-8") as f:
        f.write(corrected)

    return src_path


def parse_markdown_to_outline(md_text: str) -> list[dict]:
    """Parses strategic brief markdown into structural hierarchical outline blocks (folders, cards, text)."""
    from omniformat.core.outline import parse_markdown_to_tree

    root = parse_markdown_to_tree(md_text)
    flat_outline = []

    def traverse(node: Any) -> None:
        if node.node_type == "folder":
            flat_outline.append({"type": "folder", "title": node.title, "level": node.level})
        elif node.node_type == "card":
            flat_outline.append(
                {
                    "type": "card",
                    "tag": node.data.get("tag", ""),
                    "side": node.data.get("side", "AFF"),
                    "tier": node.data.get("tier", 1),
                    "citation": node.data.get("citation", ""),
                    "text": node.data.get("text", ""),
                    "impact_analysis": node.data.get("impact_analysis", ""),
                }
            )
        elif node.node_type == "text":
            flat_outline.append({"type": "text", "text": node.data.get("text", "")})
        for child in node.children:
            traverse(child)

    for child in root.children:
        traverse(child)
    return flat_outline


def compile_outline_to_markdown(outline: list[dict]) -> str:
    """Re-compiles structural outline blocks back into a standard, indented Markdown document."""
    from omniformat.core.outline import OutlineNode, serialize_tree_to_markdown

    root = OutlineNode("root", "STRATEGIC BRIEF: CUSTOM OUTLINE", 0)
    stack = [root]
    for item in outline:
        if item["type"] == "folder":
            level = item.get("level", 2)
            node = OutlineNode("folder", item["title"], level=level)
            while len(stack) > 1 and stack[-1].level >= level:
                stack.pop()
            stack[-1].children.append(node)
            stack.append(node)
        elif item["type"] == "card":
            node = OutlineNode("card", item.get("tag", ""), level=4, data=item)
            while len(stack) > 1 and stack[-1].level >= 4:
                stack.pop()
            stack[-1].children.append(node)
        elif item["type"] == "text":
            node = OutlineNode("text", "", level=4, data={"text": item.get("text", "")})
            while len(stack) > 1 and stack[-1].level >= 4:
                stack.pop()
            stack[-1].children.append(node)

    return serialize_tree_to_markdown(root)
